import pandas as pd
import numpy as np
import streamlit as st
import plotly.express as px
from datetime import date
from PIL import Image
from dateutil.relativedelta import relativedelta
from datetime import datetime, timedelta
from fpdf import FPDF
import tempfile
import os
import requests
import json
import copy  # NEU: Für PDF-Export hinzugefügt

######################
# Initialize variables
######################

# background nutanix logo for diagrams
background_image = dict(source=Image.open("images/nutanix-x.png"), xref="paper", yref="paper", x=0.5, y=0.5, sizex=0.95, sizey=0.95, xanchor="center", yanchor="middle", opacity=0.04, layer="below", sizing="contain")

######################
# Custom Functions
######################

# Use local CSS
def local_css(file_name):
    with open(file_name) as f:
        #st.markdown(f"<style>{f.read()}</style>", unsafe_allow_html=True)
        return f.read()

def create_date_string(numpy_datetime):
    temp_datetime = numpy_datetime.astype(datetime)
    timestring = temp_datetime.strftime("%A, %d.%m.%Y")
    return timestring

def change_input_settings():
    end_date_local = st.session_state['project_end_date']
    if (end_date_local.isoweekday() not in range(1, 6)) & (st.session_state['weekday_type']=='Arbeitstage (Mo-Fr)'):
        end_date_temp = np.busday_offset(end_date_local, 0, roll='forward')
        end_date_local = end_date_temp.astype(datetime)
        st.session_state['milestone_7_end'] = end_date_temp
    else:
        st.session_state['milestone_7_end'] = np.datetime64(end_date_local)

    if st.session_state['weekday_type'] == 'Arbeitstage (Mo-Fr)':
        st.session_state['milestone_7_start'] = np.busday_offset(st.session_state['milestone_7_end'] , -abs(st.session_state['milestone_7_duration']), roll='backward')
        st.session_state['milestone_6_start'] = np.busday_offset(st.session_state['milestone_7_start'], -abs(st.session_state['milestone_6_duration']), roll='backward')
        st.session_state['milestone_5_start'] = np.busday_offset(st.session_state['milestone_6_start'], -abs(st.session_state['milestone_5_duration']), roll='backward')
        st.session_state['milestone_4_start'] = np.busday_offset(st.session_state['milestone_5_start'], -abs(st.session_state['milestone_4_duration']), roll='backward')
        st.session_state['milestone_3_start'] = np.busday_offset(st.session_state['milestone_4_start'], -abs(st.session_state['milestone_3_duration']), roll='backward')
        st.session_state['milestone_2_start'] = np.busday_offset(st.session_state['milestone_3_start'], -abs(st.session_state['milestone_2_duration']), roll='backward')
        st.session_state['milestone_1_start'] = np.busday_offset(st.session_state['milestone_2_start'], -abs(st.session_state['milestone_1_duration']), roll='backward')

        st.session_state['milestone_6_end'] = np.busday_offset(st.session_state['milestone_7_start'], 0, roll='backward')
        st.session_state['milestone_5_end'] = np.busday_offset(st.session_state['milestone_6_start'], 0, roll='backward')
        st.session_state['milestone_4_end'] = np.busday_offset(st.session_state['milestone_5_start'], 0, roll='backward')
        st.session_state['milestone_3_end'] = np.busday_offset(st.session_state['milestone_4_start'], 0, roll='backward')
        st.session_state['milestone_2_end'] = np.busday_offset(st.session_state['milestone_3_start'], 0, roll='backward')
        st.session_state['milestone_1_end'] = np.busday_offset(st.session_state['milestone_2_start'], 0, roll='backward')
    else: # mo-so
        st.session_state['milestone_7_start'] = np.datetime64(st.session_state['milestone_7_end']) - np.timedelta64(st.session_state['milestone_7_duration'])
        st.session_state['milestone_6_start'] = np.datetime64(st.session_state['milestone_7_start']) - np.timedelta64(st.session_state['milestone_6_duration'])
        st.session_state['milestone_5_start'] = np.datetime64(st.session_state['milestone_6_start']) - np.timedelta64(st.session_state['milestone_5_duration'])
        st.session_state['milestone_4_start'] = np.datetime64(st.session_state['milestone_5_start']) - np.timedelta64(st.session_state['milestone_4_duration'])
        st.session_state['milestone_3_start'] = np.datetime64(st.session_state['milestone_4_start']) - np.timedelta64(st.session_state['milestone_3_duration'])
        st.session_state['milestone_2_start'] = np.datetime64(st.session_state['milestone_3_start']) - np.timedelta64(st.session_state['milestone_2_duration'])
        st.session_state['milestone_1_start'] = np.datetime64(st.session_state['milestone_2_start']) - np.timedelta64(st.session_state['milestone_1_duration'])

        st.session_state['milestone_6_end'] = st.session_state['milestone_7_start']
        st.session_state['milestone_5_end'] = st.session_state['milestone_6_start']
        st.session_state['milestone_4_end'] = st.session_state['milestone_5_start']
        st.session_state['milestone_3_end'] = st.session_state['milestone_4_start']
        st.session_state['milestone_2_end'] = st.session_state['milestone_3_start']
        st.session_state['milestone_1_end'] = st.session_state['milestone_2_start']

# Set Default values
def initialize_default_values():
    # Set Project end date default (today + 6 months)
    end_date = date.today() + relativedelta(months=+6)
    st.session_state['project_end_date'] = end_date
    
    # Set weekday default
    st.session_state['weekday_type'] = 'Wochentage (Mo-So)'
    
    st.session_state['milestone_1_duration'] = 21
    st.session_state['milestone_2_duration'] = 7
    st.session_state['milestone_3_duration'] = 7
    st.session_state['milestone_4_duration'] = 2
    st.session_state['milestone_5_duration'] = 28
    st.session_state['milestone_6_duration'] = 5
    st.session_state['milestone_7_duration'] = 7
    
    st.session_state['milestone_7_start'] = np.datetime64(end_date) - np.timedelta64(st.session_state['milestone_7_duration'])
    st.session_state['milestone_6_start'] = np.datetime64(st.session_state['milestone_7_start']) - np.timedelta64(st.session_state['milestone_6_duration'])
    st.session_state['milestone_5_start'] = np.datetime64(st.session_state['milestone_6_start']) - np.timedelta64(st.session_state['milestone_5_duration'])
    st.session_state['milestone_4_start'] = np.datetime64(st.session_state['milestone_5_start']) - np.timedelta64(st.session_state['milestone_4_duration'])
    st.session_state['milestone_3_start'] = np.datetime64(st.session_state['milestone_4_start']) - np.timedelta64(st.session_state['milestone_3_duration'])
    st.session_state['milestone_2_start'] = np.datetime64(st.session_state['milestone_3_start']) - np.timedelta64(st.session_state['milestone_2_duration'])
    st.session_state['milestone_1_start'] = np.datetime64(st.session_state['milestone_2_start']) - np.timedelta64(st.session_state['milestone_1_duration'])
    
    st.session_state['milestone_7_end'] = np.datetime64(st.session_state['project_end_date'])
    st.session_state['milestone_6_end'] = st.session_state['milestone_7_start']
    st.session_state['milestone_5_end'] = st.session_state['milestone_6_start']
    st.session_state['milestone_4_end'] = st.session_state['milestone_5_start']
    st.session_state['milestone_3_end'] = st.session_state['milestone_4_start']
    st.session_state['milestone_2_end'] = st.session_state['milestone_3_start']
    st.session_state['milestone_1_end'] = st.session_state['milestone_2_start']

def change_milestone_duration(milestone_number_duration):
    if st.session_state['weekday_type'] == 'Arbeitstage (Mo-Fr)':
        if milestone_number_duration >=7:
            st.session_state['milestone_7_start'] = np.busday_offset(st.session_state['milestone_7_end'], -abs(st.session_state['milestone_7_duration']), roll='backward')
        if milestone_number_duration >=6:
            st.session_state['milestone_6_start'] = np.busday_offset(st.session_state['milestone_7_start'], -abs(st.session_state['milestone_6_duration']), roll='backward')
            st.session_state['milestone_6_end'] = np.busday_offset(st.session_state['milestone_7_start'], 0, roll='backward')
        if milestone_number_duration >=5:
            st.session_state['milestone_5_start'] = np.busday_offset(st.session_state['milestone_6_start'], -abs(st.session_state['milestone_5_duration']), roll='backward')
            st.session_state['milestone_5_end'] = np.busday_offset(st.session_state['milestone_6_start'], 0, roll='backward')
        if milestone_number_duration >=4:
            st.session_state['milestone_4_start'] = np.busday_offset(st.session_state['milestone_5_start'], -abs(st.session_state['milestone_4_duration']), roll='backward')
            st.session_state['milestone_4_end'] = np.busday_offset(st.session_state['milestone_5_start'], 0, roll='backward')
        if milestone_number_duration >=3:
            st.session_state['milestone_3_start'] = np.busday_offset(st.session_state['milestone_4_start'], -abs(st.session_state['milestone_3_duration']), roll='backward')
            st.session_state['milestone_3_end'] = np.busday_offset(st.session_state['milestone_4_start'], 0, roll='backward')
        if milestone_number_duration >=2:
            st.session_state['milestone_2_start'] = np.busday_offset(st.session_state['milestone_3_start'], -abs(st.session_state['milestone_2_duration']), roll='backward')
            st.session_state['milestone_2_end'] = np.busday_offset(st.session_state['milestone_3_start'], 0, roll='backward')
        if milestone_number_duration >=1:
            st.session_state['milestone_1_start'] = np.busday_offset(st.session_state['milestone_2_start'], -abs(st.session_state['milestone_1_duration']), roll='backward')
            st.session_state['milestone_1_end'] = np.busday_offset(st.session_state['milestone_2_start'], 0, roll='backward')
    else: #mo-so
        if milestone_number_duration >=7:
            st.session_state['milestone_7_start'] = np.datetime64(st.session_state['milestone_7_end']) - np.timedelta64(st.session_state['milestone_7_duration'])
        if milestone_number_duration >=6:
            st.session_state['milestone_6_start'] = np.datetime64(st.session_state['milestone_7_start']) - np.timedelta64(st.session_state['milestone_6_duration'])
            st.session_state['milestone_6_end'] = st.session_state['milestone_7_start']
        if milestone_number_duration >=5:
            st.session_state['milestone_5_start'] = np.datetime64(st.session_state['milestone_6_start']) - np.timedelta64(st.session_state['milestone_5_duration'])
            st.session_state['milestone_5_end'] = st.session_state['milestone_6_start']
        if milestone_number_duration >=4:
            st.session_state['milestone_4_start'] = np.datetime64(st.session_state['milestone_5_start']) - np.timedelta64(st.session_state['milestone_4_duration'])
            st.session_state['milestone_4_end'] = st.session_state['milestone_5_start']
        if milestone_number_duration >=3:
            st.session_state['milestone_3_start'] = np.datetime64(st.session_state['milestone_4_start']) - np.timedelta64(st.session_state['milestone_3_duration'])
            st.session_state['milestone_3_end'] = st.session_state['milestone_4_start']
        if milestone_number_duration >=2:
            st.session_state['milestone_2_start'] = np.datetime64(st.session_state['milestone_3_start']) - np.timedelta64(st.session_state['milestone_2_duration'])
            st.session_state['milestone_2_end'] = st.session_state['milestone_3_start']
        if milestone_number_duration >=1:
            st.session_state['milestone_1_start'] = np.datetime64(st.session_state['milestone_2_start']) - np.timedelta64(st.session_state['milestone_1_duration'])
            st.session_state['milestone_1_end'] = st.session_state['milestone_2_start']

# Generate gantt Diagram
@st.cache_data  # GEÄNDERT: von @st.cache zu @st.cache_data
def generate_gantt_diagramm(gantt_df):
    gantt_diagramm = px.timeline(gantt_df,
                                  x_start="Start", 
                                  x_end="Ende",
                                  y="Meilenstein",
                                  text="Dauer"
                                  )
    
    gantt_diagramm.update_layout(
        bargap=0.5
        ,bargroupgap=0.1
        ,xaxis_range=[gantt_df.Start.min(), gantt_df.Ende.max()]
        ,xaxis = dict(
            showgrid=True
            #,rangeslider_visible=False
            #,side ="bottom"
            ,tickmode="auto"  # auto, array or linear
            ,dtick="M1"  # Set the step in-between ticks
            ,tickformat='%m.%Y'#"KW %V \n %m %Y"#"Q%q %m %Y \n"#'%m.%Y'#"%b\n%Y"#'%d.%m.%Y' #"%b\n%Y" #"Q%q %Y \n"
            ,ticklabelmode="period"#"instant"#"period"
            #,ticks="outside"
            #,tickson="boundaries"
            #,title = ''
            #,tickwidth=.1
            #,layer='below traces'
            #,ticklen=20
            #,tickfont=dict(size=20)
            ,tickfont=dict(family='sans-serif',size=16,color='black')
        ),
        yaxis = dict(
            showgrid = True,
            zeroline = True,
            showline = True,
            #gridcolor = '#bdbdbd',
            gridwidth = 2,
            zerolinecolor = '#969696',
            zerolinewidth = 2,
            linecolor = '#636363',
            linewidth = 2,
            #title = 'VALUE',
            # titlefont = dict(
            #     family = 'sans-serif',
            #     size = 18,
            #     color = 'lightgrey'
            # ),
            showticklabels = True,
            #tickangle = 45,
            tickfont = dict(
                family = 'sans-serif',
                size = 16,
                color = 'black'
            ),
            tickmode = 'linear',
            tick0 = 0.0,
            dtick = 0.25
        ))
    
    # Add vertical line for today
    today = date.today()
    gantt_diagramm.update_layout(shapes=[
        dict(
            type='line',
            yref='paper', y0=0, y1=1,
            xref='x', x0=today, x1=today)
    ])
    
    gantt_diagramm.update_yaxes(autorange="reversed")  # otherwise tasks are listed from the bottom up
    
    gantt_diagramm.update_layout(
        margin=dict(l=10, r=10, t=10, b=10,pad=4), autosize=True, #height=650,
        #xaxis={'visible': True, 'showticklabels': True}, yaxis={'visible': True, 'showticklabels': True},
        hovermode="closest",
        xaxis_title_text='',
        yaxis_title_text=''
    )
    
    gantt_diagramm.update_traces(marker=dict(color='#034EA2'))
    gantt_diagramm.update_traces(hovertemplate="<b>Start Datum:</b> %{base|%d.%m.%Y}<br>"
                                                "<b>End Datum:</b> %{x|%d.%m.%Y}<br>"
                                                "<b>Meilenstein:</b> %{y}")
    gantt_diagramm.update_traces(texttemplate='%{text} Tage', textposition='inside', insidetextanchor='middle',textfont_size=18, cliponaxis= False)
    
    gantt_diagramm_config = {
        "displaylogo": False, 'modeBarButtonsToRemove': ['zoom2d', 'toggleSpikelines', 'pan2d', 'select2d',
                                                          'lasso2d', 'autoScale2d', 'hoverClosestCartesian', 'hoverCompareCartesian']
    }
    
    gantt_diagramm.add_layout_image(background_image)
    
    return gantt_diagramm, gantt_diagramm_config

def create_excel_report(data_df, customer_name, created_by_name, gantt_diagramm, output_selection, remarks):
    """Erstellt Excel-Report mit Tabelle und optional Gantt-Diagramm"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils.dataframe import dataframe_to_rows
    from openpyxl.drawing.image import Image as XLImage
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Projektzeitraum"
    
    # Überschrift
    ws['A1'] = "Projektzeitraum Übersicht"
    ws['A1'].font = Font(size=18, bold=True, color="034EA2")
    ws.merge_cells('A1:D1')
    
    # Kundenname
    if customer_name:
        ws['A2'] = str(customer_name)
        ws['A2'].font = Font(size=12)
        ws.merge_cells('A2:D2')
        start_row = 4
    else:
        start_row = 3
    
    # Tabellen-Header
    headers = list(data_df.columns)
    header_fill = PatternFill(start_color="034EA2", end_color="034EA2", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='left')
    
    # Daten einfügen
    for row_num, row_data in enumerate(data_df.itertuples(index=False), start_row + 1):
        ws.cell(row=row_num, column=1, value=row_data[0])  # Meilenstein
        ws.cell(row=row_num, column=2, value=row_data[1])  # Dauer
        ws.cell(row=row_num, column=3, value=row_data[2].strftime("%A, %d.%m.%Y"))  # Start
        ws.cell(row=row_num, column=4, value=row_data[3].strftime("%A, %d.%m.%Y"))  # Ende
    
    # Spaltenbreiten anpassen
    ws.column_dimensions['A'].width = 40
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 25
    
    # Zusammenfassung
    summary_row = start_row + len(data_df) + 2
    if st.session_state['weekday_type'] == 'Wochentage (Mo-So)':
        summary_text = f"Der Projektzeitraum umfasst insgesamt: {data_df['Dauer'].sum()} Wochentage (Montag-Sonntag)."
    else:
        summary_text = f"Der Projektzeitraum umfasst insgesamt: {data_df['Dauer'].sum()} Arbeitstage (Montag-Freitag)."
    
    ws.cell(row=summary_row, column=1, value=summary_text)
    ws.merge_cells(f'A{summary_row}:D{summary_row}')
    
    # Anmerkungen
    if remarks:
        remarks_row = summary_row + 2
        ws.cell(row=remarks_row, column=1, value="Ergänzende Anmerkungen / Hinweise:")
        ws.cell(row=remarks_row, column=1).font = Font(bold=True, underline="single")
        ws.merge_cells(f'A{remarks_row}:D{remarks_row}')
        
        ws.cell(row=remarks_row + 1, column=1, value=remarks)
        ws.merge_cells(f'A{remarks_row + 1}:D{remarks_row + 1}')
        ws.cell(row=remarks_row + 1, column=1).alignment = Alignment(wrap_text=True)
    
    # Gantt-Diagramm einfügen (wenn gewünscht)
    if output_selection == 'Tabelle & Diagramm':
        from io import BytesIO as ImgIO

        ws2 = wb.create_sheet(title="Gantt-Diagramm")
        ws2['A1'] = "Projektzeitraum Diagramm"
        ws2['A1'].font = Font(size=18, bold=True, color="034EA2")

        gantt_export = copy.deepcopy(gantt_diagramm)
        gantt_export.update_layout(
            margin=dict(l=280, r=50, t=50, b=50),
            yaxis=dict(tickfont=dict(family='sans-serif', size=14, color='black'))
        )

        # Bild direkt im Speicher halten (kein Temp-File!)
        img_bytes = ImgIO()
        gantt_export.layout.images = []
        gantt_export.write_image(img_bytes, width=1600, height=700, format='png')
        img_bytes.seek(0)

        img = XLImage(img_bytes)
        img.width = 1200
        img.height = 525
        ws2.add_image(img, 'A3')

    # Erstellt-Informationen
    info_row = summary_row + (4 if remarks else 2)
    if created_by_name:
        ws.cell(row=info_row, column=1, value=f"Erstellt von: {created_by_name}")
        ws.cell(row=info_row + 1, column=1, value=f"Erstellt am: {date.today().strftime('%d.%m.%Y')}")
    else:
        ws.cell(row=info_row, column=1, value=f"Erstellt am: {date.today().strftime('%d.%m.%Y')}")
    
    return wb


def create_word_report(data_df, customer_name, created_by_name, gantt_diagramm, output_selection, remarks):
    """Erstellt Word-Report mit Tabelle und optional Gantt-Diagramm"""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    
    doc = Document()
    
    # Überschrift
    title = doc.add_heading('Projektzeitraum Übersicht', level=1)
    title.runs[0].font.color.rgb = RGBColor(3, 78, 162)  # Nutanix Blue
    
    # Kundenname
    if customer_name:
        customer_para = doc.add_paragraph(str(customer_name))
        customer_para.runs[0].font.size = Pt(12)
        doc.add_paragraph()  # Leerzeile
    
    # Tabelle erstellen
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Light Grid Accent 1'
    
    # Header
    header_cells = table.rows[0].cells
    headers = list(data_df.columns)
    for i, header in enumerate(headers):
        header_cells[i].text = header
        header_cells[i].paragraphs[0].runs[0].font.bold = True
        header_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(3, 78, 162)
    
    # Daten
    for _, row in data_df.iterrows():
        row_cells = table.add_row().cells
        row_cells[0].text = str(row['Meilenstein'])
        row_cells[1].text = str(row['Dauer'])
        row_cells[2].text = row['Start'].strftime("%A, %d.%m.%Y")
        row_cells[3].text = row['Ende'].strftime("%A, %d.%m.%Y")
    
    # Zusammenfassung
    doc.add_paragraph()
    if st.session_state['weekday_type'] == 'Wochentage (Mo-So)':
        summary_text = f"Der Projektzeitraum umfasst insgesamt: {data_df['Dauer'].sum()} Wochentage (Montag-Sonntag)."
    else:
        summary_text = f"Der Projektzeitraum umfasst insgesamt: {data_df['Dauer'].sum()} Arbeitstage (Montag-Freitag)."
    doc.add_paragraph(summary_text)
    
    # Anmerkungen
    if remarks:
        doc.add_paragraph()
        remarks_heading = doc.add_heading('Ergänzende Anmerkungen / Hinweise:', level=3)
        doc.add_paragraph(remarks)
    
    # Gantt-Diagramm
    if output_selection == 'Tabelle & Diagramm':
        doc.add_page_break()
        doc.add_heading('Projektzeitraum Diagramm', level=1)
        
        with tempfile.NamedTemporaryFile(delete=False, suffix=".png") as tmpfile:
            gantt_export = copy.deepcopy(gantt_diagramm)
            gantt_export.update_layout(
                margin=dict(l=280, r=50, t=50, b=50),
                yaxis=dict(tickfont=dict(family='sans-serif', size=14, color='black'))
            )
            gantt_export.layout.images = []
            gantt_export.write_image(tmpfile.name, width=1600, height=700)
            
            doc.add_picture(tmpfile.name, width=Inches(6.5))
            
            tmpfile.close()
            os.remove(tmpfile.name)
    
    # Erstellt-Informationen
    doc.add_paragraph()
    if created_by_name:
        doc.add_paragraph(f"Erstellt von: {created_by_name}")
    doc.add_paragraph(f"Erstellt am: {date.today().strftime('%d.%m.%Y')}")
    
    return doc


def create_pdf_report(data_df,customer_name,created_by_name,gantt_diagramm,output_selection,remarks):
    pdf = FPDF(format='A4', unit='mm')  # A4 (210 by 297 mm)
    pdf.add_page()
    pdf.image("./images/letterhead_cropped.png", 0, 0, 210)  # Add header picture
    pdf.set_font('Helvetica', '', 24)
    pdf.ln(40)  # line break with height
    pdf.write(4, f"Projektzeitraum Übersicht")
    pdf.ln(12)  # line break with height
    pdf.set_font('Helvetica', '', 14)
    if customer_name:
        pdf.write(3, f'{str(customer_name)}')
    pdf.ln(15)  # line break with height
    
    # header row
    pdf.set_font('Helvetica', 'B', 11)
    pdf.set_text_color(3, 78, 162)  #red,green,blue => Nutanix blue
    columnNameList = list(data_df.columns)
    pdf.cell(80, 10, columnNameList[0], 1, 0, 'L')
    pdf.cell(15, 10, columnNameList[1], 1, 0, 'L')
    pdf.cell(45, 10, columnNameList[2], 1, 0, 'L')
    pdf.cell(45, 10, columnNameList[3], 1, 0, 'L')
    pdf.ln()
    pdf.set_font('Helvetica', '', 10)
    pdf.set_text_color(0, 0, 0)  #red,green,blue => Black
    
    # milestone 1 row
    pdf.cell(80, 8, str(data_df.iloc[0,0]), 1, 0, 'L')
    pdf.cell(15, 8, str(data_df.iloc[0,1]), 1, 0, 'L')
    pdf.cell(45, 8, data_df.iloc[0,2].strftime("%A, %d.%m.%Y"), 1, 0, 'L')
    pdf.cell(45, 8, data_df.iloc[0,3].strftime("%A, %d.%m.%Y"), 1, 0, 'L')
    pdf.ln()
    
    # milestone 2 row
    pdf.cell(80, 8, str(data_df.iloc[1,0]), 1, 0, 'L')
    pdf.cell(15, 8, str(data_df.iloc[1,1]), 1, 0, 'L')
    pdf.cell(45, 8, data_df.iloc[1,2].strftime("%A, %d.%m.%Y"), 1, 0, 'L')
    pdf.cell(45, 8, data_df.iloc[1,3].strftime("%A, %d.%m.%Y"), 1, 0, 'L')
    pdf.ln()
    
    # milestone 3 row
    pdf.cell(80, 8, str(data_df.iloc[2,0]), 1, 0, 'L')
    pdf.cell(15, 8, str(data_df.iloc[2,1]), 1, 0, 'L')
    pdf.cell(45, 8, data_df.iloc[2,2].strftime("%A, %d.%m.%Y"), 1, 0, 'L')
    pdf.cell(45, 8, data_df.iloc[2,3].strftime("%A, %d.%m.%Y"), 1, 0, 'L')
    pdf.ln()
    
    # milestone 4 row
    pdf.cell(80, 8, str(data_df.iloc[3,0]), 1, 0, 'L')
    pdf.cell(15, 8, str(data_df.iloc[3,1]), 1, 0, 'L')
    pdf.cell(45, 8, data_df.iloc[3,2].strftime("%A, %d.%m.%Y"), 1, 0, 'L')
    pdf.cell(45, 8, data_df.iloc[3,3].strftime("%A, %d.%m.%Y"), 1, 0, 'L')
    pdf.ln()
    
    # milestone 5 row
    pdf.cell(80, 8, str(data_df.iloc[4,0]), 1, 0, 'L')
    pdf.cell(15, 8, str(data_df.iloc[4,1]), 1, 0, 'L')
    pdf.cell(45, 8, data_df.iloc[4,2].strftime("%A, %d.%m.%Y"), 1, 0, 'L')
    pdf.cell(45, 8, data_df.iloc[4,3].strftime("%A, %d.%m.%Y"), 1, 0, 'L')
    pdf.ln()
    
    # milestone 6 row
    pdf.cell(80, 8, str(data_df.iloc[5,0]), 1, 0, 'L')
    pdf.cell(15, 8, str(data_df.iloc[5,1]), 1, 0, 'L')
    pdf.cell(45, 8, data_df.iloc[5,2].strftime("%A, %d.%m.%Y"), 1, 0, 'L')
    pdf.cell(45, 8, data_df.iloc[5,3].strftime("%A, %d.%m.%Y"), 1, 0, 'L')
    pdf.ln()
    
    # milestone 7 row
    pdf.cell(80, 8, str(data_df.iloc[6,0]), 1, 0, 'L')
    pdf.cell(15, 8, str(data_df.iloc[6,1]), 1, 0, 'L')
    pdf.cell(45, 8, data_df.iloc[6,2].strftime("%A, %d.%m.%Y"), 1, 0, 'L')
    pdf.cell(45, 8, data_df.iloc[6,3].strftime("%A, %d.%m.%Y"), 1, 0, 'L')
    pdf.ln(15)
    
    if st.session_state['weekday_type'] == 'Wochentage (Mo-So)':
        pdf.write(3, 'Der Projektzeitraum umfasst insgesamt: '+str(data_df['Dauer'].sum())+' Wochentage (Montag-Sonntag).')
    else:
        pdf.write(3, 'Der Projektzeitraum umfasst insgesamt: '+str(data_df['Dauer'].sum())+' Arbeitstage (Montag-Freitag).')
    
    if remarks:
        pdf.ln(15)
        pdf.set_font('Helvetica', 'U', 11)
        pdf.write(5,'Ergänzende Anmerkungen / Hinweise:')
        pdf.ln(7)
        pdf.set_font('Helvetica', '', 10)
        pdf.write(5,remarks)
        pdf.ln(10)
    
    if output_selection == 'Tabelle & Diagramm':
        pdf.add_page(orientation='L')
        pdf.set_font('Helvetica', '', 24)
        pdf.ln(10)  # line break with height
        pdf.write(4, f"Projektzeitraum Diagramm")
        pdf.set_font('Helvetica', '', 10)
        pdf.ln(15)

        # Verbesserter PDF-Export mit größerem Margin
        with tempfile.NamedTemporaryFile(delete=False, suffix=".png") as tmpfile:
            # Temporäre Kopie des Diagramms erstellen
            gantt_export = copy.deepcopy(gantt_diagramm)

            # Layout für Export optimieren
            gantt_export.update_layout(
                margin=dict(l=280, r=50, t=50, b=50),  # l=280 für lange Meilenstein-Namen
                yaxis=dict(
                    tickfont=dict(
                        family='sans-serif',
                        size=14,  # Etwas kleiner für bessere Lesbarkeit
                        color='black'
                    )
                )
            )

            # Höhere Auflösung für bessere Qualität
            gantt_export.layout.images = []
            gantt_export.write_image(tmpfile.name, width=1600, height=700)

            # Bild ins PDF einfügen
            pdf.image(tmpfile.name, x=0, y=38, w=290, h=130, type='PNG', link='')

            # Temporäre Datei aufräumen
            tmpfile.close()
            os.remove(tmpfile.name)

        pdf.ln(140)

        if created_by_name:
            pdf.write(3, f'{"Erstellt von: "+str(created_by_name)}')
            pdf.ln(5)
            pdf.write(3, f'{"Erstellt am: "+date.today().strftime("%d.%m.%Y")}')
        else:
            pdf.ln(5)
            pdf.write(3, f'{"Erstellt am: "+date.today().strftime("%d.%m.%Y")}')

    return pdf


# Send Slack Message
# NO cache function!
def send_slack_message():
    # Send a Slack message to a channel via a webhook.
    webhook = st.secrets["slack_webhook_url"]
    payload = {"text": 'Reverse Timeline Planning - Report erstellt.'}
    requests.post(webhook, json.dumps(payload))
